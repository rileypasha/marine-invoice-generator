"""
Read VesselsandContactsCombined.xlsx and emit a JSON array of rows.
Usage: python xlsx-to-json.py <input.xlsx> <output.json>
"""
import json
import sys
import openpyxl


def normalize(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def main(inp: str, out: str) -> None:
    wb = openpyxl.load_workbook(inp, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [normalize(c.value) or f"col_{i}" for i, c in enumerate(ws[1])]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {headers[i]: normalize(v) for i, v in enumerate(row) if i < len(headers)}
        # skip totally empty rows
        if any(record.values()):
            rows.append(record)
    with open(out, "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2, ensure_ascii=False)
    print(f"Wrote {len(rows)} rows to {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
